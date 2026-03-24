"""
导入甪直站 Excel SCADA 历史数据到 scada_history 表

Excel 列结构（Sheet1）:
  列 0-4:   二线水露点 (Date, Time, Value, -, -)
  列 5-8:   中俄水露点 (Date, Time, Value, -)
  列 9-13:  西一线压力 (Date, Time, Value, -, -)
  列 14-18: 西一线温度 (Date, Time, Value, -, -)
  列 19-23: 西二线压力 (Date, Time, Value, -, -)
  列 24-28: 西二线温度 (Date, Time, Value, -, -)
  列 29-32: 中俄线压力 (Date, Time, Value, -)
  列 33-35: 中俄线温度 (Date, Time, Value)
"""
import sys
sys.stdout.reconfigure(encoding='utf-8')

import openpyxl
from datetime import datetime, timedelta
from pathlib import Path

# 添加项目路径
sys.path.insert(0, str(Path(__file__).parent))

from app.database import engine
from app.scada_models import ScadaHistory
from sqlmodel import Session, text

EXCEL_PATH = r"C:\Users\jushixio\Downloads\甪直站20262311-0312.xlsx"
STATION_NAME = "甪直分输站"

# 定义列映射: (date_col, time_col, value_col, pipeline_id, metric_type, tag_name)
COLUMN_MAPPINGS = [
    (9,  10, 11, "we1",  "pressure",    "XDX_02C006_PI1201.PV"),   # 西一线压力
    (14, 15, 16, "we1",  "temperature", "XDX_02C006_TI1201.PV"),   # 西一线温度
    (19, 20, 21, "we2",  "pressure",    "XDE_02C006_PT1102.PV"),   # 西二线压力
    (24, 25, 26, "we2",  "temperature", "XDE_02C006_TT1101.PV"),   # 西二线温度
    (29, 30, 31, "cred", "pressure",    "ZE_PT1101.PV"),           # 中俄线压力
    (33, 34, 35, "cred", "temperature", "ZE_TT1201.PV"),           # 中俄线温度
]

def combine_datetime(date_val, time_val):
    """合并日期和时间列为 datetime 对象"""
    if date_val is None or time_val is None:
        return None
    
    if isinstance(date_val, datetime):
        base_date = date_val
    else:
        return None
    
    from datetime import time as time_type
    if isinstance(time_val, time_type):
        return datetime(base_date.year, base_date.month, base_date.day,
                       time_val.hour, time_val.minute, time_val.second)
    
    return None

def main():
    print(f"正在读取 Excel: {EXCEL_PATH}")
    wb = openpyxl.load_workbook(EXCEL_PATH, read_only=True)
    ws = wb[wb.sheetnames[0]]
    
    # 跳过前 2 行表头
    rows = list(ws.iter_rows(min_row=3, values_only=True))
    print(f"数据行数: {len(rows)}")
    
    records = []
    for row in rows:
        for date_col, time_col, val_col, pipeline_id, metric_type, tag_name in COLUMN_MAPPINGS:
            try:
                date_val = row[date_col]
                time_val = row[time_col]
                value = row[val_col]
                
                if value is None or date_val is None or time_val is None:
                    continue
                
                recorded_at = combine_datetime(date_val, time_val)
                if recorded_at is None:
                    continue
                
                records.append(ScadaHistory(
                    station_name=STATION_NAME,
                    pipeline_id=pipeline_id,
                    tag_name=tag_name,
                    metric_type=metric_type,
                    value=float(value),
                    recorded_at=recorded_at,
                ))
            except (IndexError, TypeError, ValueError) as e:
                continue
    
    print(f"解析到 {len(records)} 条有效记录")
    
    # 写入数据库
    with Session(engine) as session:
        # 先清除该站的旧数据
        session.exec(text(f"DELETE FROM scada_history WHERE station_name = '{STATION_NAME}'"))
        session.commit()
        
        # 批量插入
        for r in records:
            session.add(r)
        session.commit()
        
        # 验证
        from sqlmodel import select, func
        count = session.exec(
            select(func.count()).where(ScadaHistory.station_name == STATION_NAME)
        ).one()
        print(f"已导入 {count} 条记录到 scada_history 表")
        
        # 按指标统计
        for pipeline_id, metric_type in [("we1", "pressure"), ("we1", "temperature"), 
                                          ("we2", "pressure"), ("we2", "temperature"),
                                          ("cred", "pressure"), ("cred", "temperature")]:
            c = session.exec(
                select(func.count()).where(
                    ScadaHistory.station_name == STATION_NAME,
                    ScadaHistory.pipeline_id == pipeline_id,
                    ScadaHistory.metric_type == metric_type,
                )
            ).one()
            print(f"  {pipeline_id}_{metric_type}: {c} 条")

if __name__ == "__main__":
    main()
