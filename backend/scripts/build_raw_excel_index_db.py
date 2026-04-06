from __future__ import annotations

import argparse
import json
import re
import sqlite3
from dataclasses import dataclass
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

import openpyxl


ROOT_DIR = Path(__file__).resolve().parents[2]
BACKEND_DIR = ROOT_DIR / "backend"
DATA_DIR = BACKEND_DIR / "data"
RAW_XLSX_DIR = DATA_DIR / "raw_csvs"
DEFAULT_DB_PATH = DATA_DIR / "raw_excel_index.db"


@dataclass
class SheetPlan:
    sheet_name: str
    table_name: str
    header_row_index: int
    headers: list[str]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Build a raw SQLite index database from an Excel workbook.")
    parser.add_argument(
        "--xlsx",
        type=Path,
        default=None,
        help="Path to the source xlsx. Defaults to the first usable file under backend/data/raw_csvs.",
    )
    parser.add_argument(
        "--db",
        type=Path,
        default=DEFAULT_DB_PATH,
        help="Output SQLite path. Default: backend/data/raw_excel_index.db",
    )
    parser.add_argument(
        "--replace",
        action="store_true",
        help="Delete the target db first if it already exists.",
    )
    return parser.parse_args()


def resolve_xlsx_path(explicit_path: Path | None) -> Path:
    if explicit_path:
        return explicit_path.resolve()

    files = sorted(path for path in RAW_XLSX_DIR.glob("*.xlsx") if not path.name.startswith("~$"))
    if not files:
        raise FileNotFoundError(f"No usable xlsx found under {RAW_XLSX_DIR}")
    return files[0]


def ensure_parent_dir(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)


def sanitize_identifier(raw_name: str, fallback: str) -> str:
    text = re.sub(r"\s+", "_", raw_name.strip())
    text = re.sub(r"[^\w\u4e00-\u9fff]+", "_", text)
    text = text.strip("_")
    if not text:
        text = fallback
    if text[0].isdigit():
        text = f"c_{text}"
    return text.lower()


def dedupe_headers(raw_headers: list[str]) -> list[str]:
    seen: dict[str, int] = {}
    normalized_headers: list[str] = []

    for index, header in enumerate(raw_headers, start=1):
        base_name = sanitize_identifier(header, f"col_{index}")
        count = seen.get(base_name, 0) + 1
        seen[base_name] = count
        normalized_headers.append(base_name if count == 1 else f"{base_name}_{count}")

    return normalized_headers


def to_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(sep=" ", timespec="seconds")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.isoformat(timespec="seconds")
    if isinstance(value, bool):
        return "true" if value else "false"
    return str(value).strip()


def is_text_like(value: Any) -> bool:
    text = to_text(value)
    if not text:
        return False
    if re.fullmatch(r"[-+]?\d+(?:\.\d+)?", text):
        return False
    return True


def choose_header_row(preview_rows: list[tuple[Any, ...]]) -> int:
    best_index = 0
    best_score = float("-inf")

    for index, row in enumerate(preview_rows):
        cells = [to_text(cell) for cell in row]
        non_empty = [cell for cell in cells if cell]
        if len(non_empty) < 2:
            continue

        text_count = sum(1 for cell in row if is_text_like(cell))
        numeric_count = sum(1 for cell in non_empty if re.fullmatch(r"[-+]?\d+(?:\.\d+)?", cell))
        duplicate_penalty = len(non_empty) - len(set(non_empty))
        score = (text_count * 3) + len(non_empty) - (numeric_count * 2) - duplicate_penalty - (index * 0.05)

        if score > best_score:
            best_score = score
            best_index = index

    return best_index + 1


def build_sheet_plan(sheet_name: str, worksheet: openpyxl.worksheet.worksheet.Worksheet) -> SheetPlan:
    preview_rows = list(worksheet.iter_rows(min_row=1, max_row=min(25, worksheet.max_row), values_only=True))
    header_row_index = choose_header_row(preview_rows)
    header_row = preview_rows[header_row_index - 1]
    raw_headers = [to_text(cell).replace("\n", " ").strip() for cell in header_row]
    headers = dedupe_headers(raw_headers)
    table_name = sanitize_identifier(sheet_name, "sheet")
    return SheetPlan(
        sheet_name=sheet_name,
        table_name=f"sheet_{table_name}",
        header_row_index=header_row_index,
        headers=headers,
    )


def create_base_schema(conn: sqlite3.Connection) -> None:
    conn.executescript(
        """
        PRAGMA journal_mode=WAL;
        PRAGMA synchronous=NORMAL;

        CREATE TABLE IF NOT EXISTS workbook_files (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            file_name TEXT NOT NULL,
            file_path TEXT NOT NULL,
            imported_at TEXT NOT NULL,
            file_size_bytes INTEGER NOT NULL,
            sheet_count INTEGER NOT NULL
        );

        CREATE TABLE IF NOT EXISTS workbook_sheets (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            workbook_id INTEGER NOT NULL,
            sheet_name TEXT NOT NULL,
            table_name TEXT NOT NULL,
            header_row_index INTEGER NOT NULL,
            row_count INTEGER NOT NULL DEFAULT 0,
            column_count INTEGER NOT NULL DEFAULT 0,
            FOREIGN KEY (workbook_id) REFERENCES workbook_files(id)
        );

        CREATE TABLE IF NOT EXISTS workbook_sheet_columns (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            sheet_id INTEGER NOT NULL,
            column_ordinal INTEGER NOT NULL,
            source_header TEXT NOT NULL,
            column_name TEXT NOT NULL,
            FOREIGN KEY (sheet_id) REFERENCES workbook_sheets(id)
        );

        CREATE TABLE IF NOT EXISTS indexed_rows (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            workbook_id INTEGER NOT NULL,
            sheet_id INTEGER NOT NULL,
            sheet_name TEXT NOT NULL,
            table_name TEXT NOT NULL,
            row_number INTEGER NOT NULL,
            row_json TEXT NOT NULL,
            row_text TEXT NOT NULL,
            FOREIGN KEY (workbook_id) REFERENCES workbook_files(id),
            FOREIGN KEY (sheet_id) REFERENCES workbook_sheets(id)
        );

        CREATE INDEX IF NOT EXISTS idx_indexed_rows_sheet ON indexed_rows(sheet_name, row_number);
        CREATE INDEX IF NOT EXISTS idx_indexed_rows_table ON indexed_rows(table_name, row_number);
        """
    )

    try:
        conn.execute(
            """
            CREATE VIRTUAL TABLE IF NOT EXISTS indexed_rows_fts
            USING fts5(sheet_name, table_name, row_text)
            """
        )
    except sqlite3.OperationalError:
        pass


def create_sheet_table(conn: sqlite3.Connection, table_name: str, headers: list[str]) -> None:
    quoted_columns = ",\n".join(f'"{header}" TEXT' for header in headers)
    conn.execute(f'DROP TABLE IF EXISTS "{table_name}"')
    conn.execute(
        f"""
        CREATE TABLE "{table_name}" (
            "_row_number" INTEGER NOT NULL,
            {quoted_columns}
        )
        """
    )
    conn.execute(f'CREATE INDEX "idx_{table_name}_row_number" ON "{table_name}"("_row_number")')


def insert_workbook_file(conn: sqlite3.Connection, xlsx_path: Path, sheet_count: int) -> int:
    cursor = conn.execute(
        """
        INSERT INTO workbook_files (file_name, file_path, imported_at, file_size_bytes, sheet_count)
        VALUES (?, ?, ?, ?, ?)
        """,
        (
            xlsx_path.name,
            str(xlsx_path),
            datetime.now().isoformat(timespec="seconds"),
            xlsx_path.stat().st_size,
            sheet_count,
        ),
    )
    return int(cursor.lastrowid)


def insert_sheet_meta(conn: sqlite3.Connection, workbook_id: int, plan: SheetPlan) -> int:
    cursor = conn.execute(
        """
        INSERT INTO workbook_sheets (workbook_id, sheet_name, table_name, header_row_index, row_count, column_count)
        VALUES (?, ?, ?, ?, 0, ?)
        """,
        (workbook_id, plan.sheet_name, plan.table_name, plan.header_row_index, len(plan.headers)),
    )
    return int(cursor.lastrowid)


def insert_sheet_columns(conn: sqlite3.Connection, sheet_id: int, source_headers: list[str], headers: list[str]) -> None:
    rows = [
        (sheet_id, ordinal, source_header, column_name)
        for ordinal, (source_header, column_name) in enumerate(zip(source_headers, headers), start=1)
    ]
    conn.executemany(
        """
        INSERT INTO workbook_sheet_columns (sheet_id, column_ordinal, source_header, column_name)
        VALUES (?, ?, ?, ?)
        """,
        rows,
    )


def update_sheet_stats(conn: sqlite3.Connection, sheet_id: int, row_count: int) -> None:
    conn.execute("UPDATE workbook_sheets SET row_count = ? WHERE id = ?", (row_count, sheet_id))


def quote_identifier(name: str) -> str:
    return '"' + name.replace('"', '""') + '"'


def sync_fts(conn: sqlite3.Connection) -> None:
    try:
        conn.execute("DELETE FROM indexed_rows_fts")
        conn.execute(
            """
            INSERT INTO indexed_rows_fts(rowid, sheet_name, table_name, row_text)
            SELECT id, sheet_name, table_name, row_text
            FROM indexed_rows
            """
        )
    except sqlite3.OperationalError:
        pass


def build_database(xlsx_path: Path, db_path: Path, replace: bool) -> dict[str, Any]:
    if replace and db_path.exists():
        db_path.unlink()

    ensure_parent_dir(db_path)

    workbook = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    conn = sqlite3.connect(db_path)
    try:
        create_base_schema(conn)
        workbook_id = insert_workbook_file(conn, xlsx_path, len(workbook.sheetnames))

        summary: dict[str, Any] = {
            "xlsx": str(xlsx_path),
            "db": str(db_path),
            "sheet_count": len(workbook.sheetnames),
            "sheets": [],
        }

        for sheet_name in workbook.sheetnames:
            worksheet = workbook[sheet_name]
            plan = build_sheet_plan(sheet_name, worksheet)
            raw_header_row = list(
                worksheet.iter_rows(
                    min_row=plan.header_row_index,
                    max_row=plan.header_row_index,
                    values_only=True,
                )
            )[0]
            source_headers = [to_text(cell).replace("\n", " ").strip() for cell in raw_header_row]

            create_sheet_table(conn, plan.table_name, plan.headers)
            sheet_id = insert_sheet_meta(conn, workbook_id, plan)
            insert_sheet_columns(conn, sheet_id, source_headers, plan.headers)

            row_count = 0
            insert_columns = ['"_row_number"'] + [quote_identifier(header) for header in plan.headers]
            placeholders = ", ".join(["?"] * len(insert_columns))
            insert_sql = f'INSERT INTO "{plan.table_name}" ({", ".join(insert_columns)}) VALUES ({placeholders})'

            for row_number, row in enumerate(
                worksheet.iter_rows(min_row=plan.header_row_index + 1, values_only=True),
                start=plan.header_row_index + 1,
            ):
                values = [to_text(cell) for cell in row[: len(plan.headers)]]
                if not any(values):
                    continue

                if len(values) < len(plan.headers):
                    values.extend([""] * (len(plan.headers) - len(values)))

                conn.execute(insert_sql, [row_number, *values])

                row_payload = {header: value for header, value in zip(plan.headers, values)}
                row_text = " | ".join(value for value in values if value)
                conn.execute(
                    """
                    INSERT INTO indexed_rows (workbook_id, sheet_id, sheet_name, table_name, row_number, row_json, row_text)
                    VALUES (?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        workbook_id,
                        sheet_id,
                        plan.sheet_name,
                        plan.table_name,
                        row_number,
                        json.dumps(row_payload, ensure_ascii=False),
                        row_text,
                    ),
                )
                row_count += 1

            update_sheet_stats(conn, sheet_id, row_count)
            summary["sheets"].append(
                {
                    "sheet_name": plan.sheet_name,
                    "table_name": plan.table_name,
                    "header_row_index": plan.header_row_index,
                    "row_count": row_count,
                    "column_count": len(plan.headers),
                }
            )

        sync_fts(conn)
        conn.commit()
        return summary
    finally:
        workbook.close()
        conn.close()


def print_summary(summary: dict[str, Any]) -> None:
    print("=" * 72)
    print("Raw Excel index database created")
    print("xlsx:", summary["xlsx"])
    print("db :", summary["db"])
    print("sheets:", summary["sheet_count"])
    print("-" * 72)
    for sheet in summary["sheets"]:
        print(
            f'{sheet["sheet_name"]} -> {sheet["table_name"]} | '
            f'header_row={sheet["header_row_index"]} | '
            f'rows={sheet["row_count"]} | cols={sheet["column_count"]}'
        )


def main() -> None:
    args = parse_args()
    xlsx_path = resolve_xlsx_path(args.xlsx)
    summary = build_database(xlsx_path=xlsx_path, db_path=args.db.resolve(), replace=args.replace)
    print_summary(summary)


if __name__ == "__main__":
    main()
