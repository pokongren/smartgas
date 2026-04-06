"""
解析甪直站 Excel 历史数据，转换为前端可用的 JSON 格式
"""
import openpyxl
import json
from datetime import datetime, time

wb = openpyxl.load_workbook(r'C:\Users\Administrator\Downloads\甪直站20262311-0312.xlsx', data_only=True)
ws = wb.active

# 列映射：指标名 -> (日期列, 时间列, 数值列)
col_map = {
    '西一线压力': (10, 11, 12),
    '西一线温度': (15, 16, 17),
    '西二线压力': (20, 21, 22),
    '西二线温度': (25, 26, 27),
    '中俄线压力': (30, 31, 32),
    '中俄线温度': (34, 35, 36),
    '二线水露点': (1,  2,  3),
    '中俄水露点': (6,  7,  8),
}

result = {}
for metric, (dc, tc, vc) in col_map.items():
    records = []
    for r in range(3, ws.max_row + 1):
        d = ws.cell(r, dc).value
        t = ws.cell(r, tc).value
        v = ws.cell(r, vc).value
        if d and t and v is not None:
            date_str = d.strftime('%Y-%m-%d') if isinstance(d, datetime) else str(d)
            time_str = t.strftime('%H:%M') if isinstance(t, time) else str(t)
            records.append({'ts': date_str + ' ' + time_str, 'v': round(float(v), 3)})
    result[metric] = records
    print(f"{metric}: {len(records)} 条, {records[0]['ts']} ~ {records[-1]['ts']}, 样本={[r['v'] for r in records[:3]]}")

# 输出 JSON
out_path = r'f:\smartgas-grid\src\data\luzhi_history.json'
with open(out_path, 'w', encoding='utf-8') as f:
    json.dump(result, f, ensure_ascii=False, indent=2)

print(f"\n已输出到: {out_path}")
print(f"总指标数: {len(result)}")
for k, v in result.items():
    print(f"  {k}: {len(v)} 条记录")
